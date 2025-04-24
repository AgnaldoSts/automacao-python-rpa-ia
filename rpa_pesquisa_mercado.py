import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from datetime import datetime
import smtplib
import ssl
from email.message import EmailMessage
import tkinter as tk
from tkinter import messagebox, filedialog

# Função principal da automação
def executar_automacao(pesquisa, limite_resultados, email_destino, senha_app):
    print("🔄 Iniciando automação no Bing...")

    chrome_options = Options()
    chrome_options.add_argument("--start-maximized")
    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=chrome_options)
    driver.get("https://www.bing.com")

    campo_busca = driver.find_element(By.NAME, "q")
    campo_busca.send_keys(pesquisa)
    campo_busca.submit()

    wait = WebDriverWait(driver, 10)
    try:
        itens = wait.until(EC.presence_of_all_elements_located((By.CSS_SELECTOR, "li.b_algo")))[:limite_resultados]
    except:
        print("❌ Não foi possível carregar os resultados.")
        driver.quit()
        return

    resultados = []
    for idx, item in enumerate(itens, 1):
        try:
            titulo = item.find_element(By.TAG_NAME, "h2").text
            link = item.find_element(By.TAG_NAME, "a").get_attribute("href")
            descricao = item.find_element(By.CLASS_NAME, "b_caption").text
        except:
            continue

        resultados.append({
            "Posição": idx,
            "Título": titulo,
            "URL": link,
            "Descrição": descricao
        })

    driver.quit()

    df = pd.DataFrame(resultados)
    print("📊 Resultados da pesquisa (Bing):")
    print(df[["Posição", "Título"]].to_string(index=False))

    # Exportar para Excel
    data_hora = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    nome_arquivo = f"resultados_bing_{data_hora}.xlsx"
    with pd.ExcelWriter(nome_arquivo, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Resultados", index=False)

    # Gráfico
    wb = load_workbook(nome_arquivo)
    ws = wb.active
    contagem_titulos = df["Título"].value_counts().reset_index()
    contagem_titulos.columns = ["Título", "Ocorrências"]
    for row in contagem_titulos.values.tolist():
        ws.append(row)

    chart = BarChart()
    chart.title = "Frequência de Títulos"
    chart.add_data(Reference(ws, min_col=2, min_row=limite_resultados+3, max_row=limite_resultados+2+len(contagem_titulos)), titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=1, min_row=limite_resultados+3, max_row=limite_resultados+2+len(contagem_titulos)))
    ws.add_chart(chart, f"E5")
    wb.save(nome_arquivo)

    print(f"✅ Relatório salvo como: {nome_arquivo}")

    # Enviar por e-mail
    try:
        enviar_email(nome_arquivo, email_destino, senha_app)
        print("📧 E-mail enviado com sucesso!")
    except Exception as e:
        print(f"❌ Falha ao enviar e-mail: {e}")

    messagebox.showinfo("Finalizado", "Automação concluída com sucesso!")

def enviar_email(arquivo_excel, destino, senha):
    email_origem = destino
    assunto = "Relatório de Pesquisa - RPA Python"
    corpo = "Segue em anexo o relatório da pesquisa automatizada no Bing."

    msg = EmailMessage()
    msg["From"] = email_origem
    msg["To"] = destino
    msg["Subject"] = assunto
    msg.set_content(corpo)

    with open(arquivo_excel, "rb") as f:
        conteudo = f.read()
        msg.add_attachment(conteudo, maintype="application", subtype="xlsx", filename=arquivo_excel)

    contexto = ssl.create_default_context()
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=contexto) as smtp:
        smtp.login(email_origem, senha)
        smtp.send_message(msg)

# Interface gráfica (Tkinter)
def iniciar_interface():
    def ao_clicar():
        pesquisa = entrada_pesquisa.get()
        limite = int(entrada_limite.get())
        email = entrada_email.get()
        senha = entrada_senha.get()

        if not (pesquisa and email and senha):
            messagebox.showerror("Erro", "Todos os campos são obrigatórios.")
            return

        executar_automacao(pesquisa, limite, email, senha)

    janela = tk.Tk()
    janela.title("Pesquisa de Mercado - RPA com IA")
    janela.geometry("400x300")

    tk.Label(janela, text="🔍 Palavra-chave:").pack(pady=5)
    entrada_pesquisa = tk.Entry(janela, width=40)
    entrada_pesquisa.pack()

    tk.Label(janela, text="📈 Limite de resultados:").pack(pady=5)
    entrada_limite = tk.Entry(janela, width=10)
    entrada_limite.insert(0, "10")
    entrada_limite.pack()

    tk.Label(janela, text="📧 Seu e-mail (Gmail):").pack(pady=5)
    entrada_email = tk.Entry(janela, width=40)
    entrada_email.pack()

    tk.Label(janela, text="🔑 Senha do app Gmail:").pack(pady=5)
    entrada_senha = tk.Entry(janela, show="*", width=40)
    entrada_senha.pack()

    tk.Button(janela, text="🚀 Iniciar Automação", command=ao_clicar, bg="#28a745", fg="white").pack(pady=20)
    janela.mainloop()

# Rodar
if __name__ == "__main__":
    iniciar_interface()



